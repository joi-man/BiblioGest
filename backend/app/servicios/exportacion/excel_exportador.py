"""Convierte una TablaExportable en un libro de Excel (.xlsx)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.servicios.exportacion.tabla_exportable import TablaExportable

_COLOR_ENCABEZADO = "3B2F2A"
_ANCHO_MINIMO = 10
_ANCHO_MAXIMO = 50


def a_excel(tabla: TablaExportable) -> bytes:
    """Genera el archivo .xlsx con encabezado destacado y columnas autoajustadas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Listado"

    hoja.append([tabla.titulo])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append([tabla.descripcion])
    hoja.append([])  # fila en blanco antes de la tabla

    fila_encabezado = hoja.max_row + 1
    hoja.append(tabla.columnas)
    for celda in hoja[fila_encabezado]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=_COLOR_ENCABEZADO)
        celda.alignment = Alignment(vertical="center")

    for fila in tabla.filas:
        hoja.append(fila)
        _forzar_texto(hoja[hoja.max_row])

    _ajustar_anchos(hoja, tabla, fila_encabezado)
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _forzar_texto(celdas) -> None:
    """Marca cada celda como texto.

    Evita que un valor escrito por un usuario que empiece con "=" se ejecute
    como formula al abrir el archivo.
    """
    for celda in celdas:
        if isinstance(celda.value, str):
            celda.data_type = "s"


def _ajustar_anchos(hoja, tabla: TablaExportable, fila_encabezado: int) -> None:
    """Ancho de cada columna segun su texto mas largo, con minimo y maximo."""
    for indice, columna in enumerate(tabla.columnas, start=1):
        mas_largo = max(
            [len(columna)] + [len(str(fila[indice - 1])) for fila in tabla.filas]
        )
        ancho = min(max(mas_largo + 2, _ANCHO_MINIMO), _ANCHO_MAXIMO)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
