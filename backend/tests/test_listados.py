"""Pruebas de busqueda paginada y exportacion de usuarios y materiales."""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import _crear_usuario, encabezados


def _crear_muchos_usuarios(db, cantidad: int) -> None:
    for i in range(cantidad):
        _crear_usuario(db, f"alumno{i:02d}", "estudiante", f"E-{i:03d}")


# ------------------------------------------------------------------ usuarios
def test_usuarios_paginados_devuelven_total_y_paginas(cliente, db, admin):
    _crear_muchos_usuarios(db, 25)  # mas el admin = 26

    r = cliente.get("/api/usuarios/buscar", params={"pagina": 3, "tamano": 10}, headers=encabezados(admin))

    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["total"] == 26
    assert cuerpo["total_paginas"] == 3
    assert cuerpo["pagina"] == 3
    assert len(cuerpo["items"]) == 6


def test_usuarios_paginas_no_repiten_filas(cliente, db, admin):
    _crear_muchos_usuarios(db, 14)

    ids = []
    for pagina in (1, 2):
        r = cliente.get("/api/usuarios/buscar", params={"pagina": pagina, "tamano": 8}, headers=encabezados(admin))
        ids += [u["id"] for u in r.json()["items"]]

    assert len(ids) == len(set(ids)) == 15


def test_usuarios_busca_por_nombre_completo_y_filtra_por_rol(cliente, db, admin, estudiante):
    cabeceras = encabezados(admin)

    por_nombre = cliente.get("/api/usuarios/buscar", params={"q": "Ana Prueba"}, headers=cabeceras).json()
    por_rol = cliente.get("/api/usuarios/buscar", params={"rol": "admin"}, headers=cabeceras).json()

    assert [u["username"] for u in por_nombre["items"]] == ["ana"]
    assert [u["username"] for u in por_rol["items"]] == ["admin"]


def test_usuarios_filtra_por_estado(cliente, db, admin, estudiante):
    estudiante.activo = False
    db.commit()

    r = cliente.get("/api/usuarios/buscar", params={"activo": False}, headers=encabezados(admin))

    assert [u["username"] for u in r.json()["items"]] == ["ana"]


def test_tamano_de_pagina_tiene_tope(cliente, admin):
    r = cliente.get("/api/usuarios/buscar", params={"tamano": 1000}, headers=encabezados(admin))

    assert r.status_code == 422


def test_estudiante_no_puede_buscar_ni_exportar_usuarios(cliente, estudiante):
    cabeceras = encabezados(estudiante)

    assert cliente.get("/api/usuarios/buscar", headers=cabeceras).status_code == 403
    assert cliente.get("/api/usuarios/exportar", headers=cabeceras).status_code == 403


def test_usuarios_exporta_excel_con_todas_las_filas_filtradas(cliente, db, admin):
    _crear_muchos_usuarios(db, 15)

    r = cliente.get("/api/usuarios/exportar", params={"formato": "xlsx", "rol": "estudiante"}, headers=encabezados(admin))

    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "usuarios_" in r.headers["content-disposition"]
    hoja = load_workbook(BytesIO(r.content)).active
    filas_datos = [fila for fila in hoja.iter_rows(min_row=5, values_only=True) if fila[0]]
    assert len(filas_datos) == 15  # no solo una pagina


def test_usuarios_exporta_pdf(cliente, admin):
    r = cliente.get("/api/usuarios/exportar", params={"formato": "pdf"}, headers=encabezados(admin))

    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")


def test_exportar_formato_desconocido_da_400(cliente, admin):
    r = cliente.get("/api/usuarios/exportar", params={"formato": "doc"}, headers=encabezados(admin))

    assert r.status_code == 400


def test_excel_no_ejecuta_formulas_escritas_por_usuarios(cliente, db, admin):
    malicioso = _crear_usuario(db, "mal", "estudiante", "X-1")
    malicioso.nombre = "=1+1"
    db.commit()

    r = cliente.get("/api/usuarios/exportar", params={"formato": "xlsx", "q": "mal"}, headers=encabezados(admin))

    hoja = load_workbook(BytesIO(r.content)).active
    celda = next(c for fila in hoja.iter_rows(min_row=5) for c in fila if c.value == "=1+1")
    assert celda.data_type == "s"  # texto, no formula


# --------------------------------------------------------------- materiales
def test_materiales_paginados_y_filtrados(cliente, db, admin, categoria_libros, categoria_equipos, crear_material):
    for i in range(12):
        crear_material(categoria_libros, f"L-{i:02d}", titulo=f"Libro {i}")
    crear_material(categoria_equipos, "P-01", titulo="Proyector", tipo_material="proyector")
    cabeceras = encabezados(admin)

    pagina = cliente.get("/api/materiales/buscar", params={"tamano": 5, "pagina": 3}, headers=cabeceras).json()
    proyectores = cliente.get("/api/materiales/buscar", params={"tipo": "proyector"}, headers=cabeceras).json()
    por_categoria = cliente.get(
        "/api/materiales/buscar", params={"categoria_id": str(categoria_libros.id)}, headers=cabeceras
    ).json()

    assert pagina["total"] == 13 and pagina["total_paginas"] == 3 and len(pagina["items"]) == 3
    assert [m["codigo"] for m in proyectores["items"]] == ["P-01"]
    assert por_categoria["total"] == 12


def test_materiales_exporta_excel_con_categoria(cliente, admin, categoria_libros, crear_material):
    crear_material(categoria_libros, "L-01", titulo="Rayuela")

    r = cliente.get("/api/materiales/exportar", params={"formato": "xlsx"}, headers=encabezados(admin))

    assert r.status_code == 200
    hoja = load_workbook(BytesIO(r.content)).active
    fila = [c for c in next(hoja.iter_rows(min_row=5, max_row=5, values_only=True))]
    assert fila[0] == "L-01" and fila[3] == "Libros"


def test_materiales_exporta_pdf(cliente, admin, categoria_libros, crear_material):
    crear_material(categoria_libros, "L-01")

    r = cliente.get("/api/materiales/exportar", params={"formato": "pdf"}, headers=encabezados(admin))

    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_estudiante_no_puede_buscar_ni_exportar_materiales_de_gestion(cliente, estudiante):
    cabeceras = encabezados(estudiante)

    assert cliente.get("/api/materiales/buscar", headers=cabeceras).status_code == 403
    assert cliente.get("/api/materiales/exportar", headers=cabeceras).status_code == 403
    assert cliente.get("/api/materiales/", headers=cabeceras).status_code == 200  # catalogo sigue abierto
